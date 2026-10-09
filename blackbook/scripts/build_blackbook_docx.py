import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = r"c:\Users\Aditya\Desktop\cafe2\blackbook"
fig_dir = os.path.join(base_dir, "figures")
docx_path = os.path.join(base_dir, "BLACKBOOK.docx")

doc = Document()

# Configure Margins: Left 1.5 inches (Binding), Right 1.0 in, Top 1.0 in, Bottom 1.0 in
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = False

# Configure Normal style
style_normal = doc.styles['Normal']
font = style_normal.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0x2E, 0x1F, 0x18)
style_normal.paragraph_format.line_spacing = 1.5
style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
style_normal.paragraph_format.space_after = Pt(6)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_chapter_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x2E, 0x1F, 0x18)
    return p

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x2E, 0x1F, 0x18)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x3B, 0x2E, 0x28)
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x5C, 0x4A, 0x3E)
    return p

def add_body_paragraph(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.alignment = align
    p.paragraph_format.space_after = Pt(6)
    
    # Simple markdown parser for bold, italics, code
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            p.add_run(token)
    return p

def add_bullet_item(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            p.add_run(token)
    return p

def add_code_block(code_lines, caption=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F6F2")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.line_spacing = 1.15
    cp.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(code_lines):
        if i > 0:
            cp = cell.add_paragraph()
            cp.paragraph_format.line_spacing = 1.15
            cp.paragraph_format.space_after = Pt(0)
        run = cp.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x24, 0x1A, 0x14)
        
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        rc = p_cap.add_run(caption)
        rc.font.size = Pt(9.5)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(0x5C, 0x4A, 0x3E)

def add_figure_image(image_filename, caption_text, width=Inches(5.8)):
    img_path = os.path.join(fig_dir, image_filename)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=width)
        
        # Caption
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(14)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10)
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(0x2E, 0x1F, 0x18)
    else:
        print(f"Warning: Figure image not found at {img_path}")

def parse_markdown_table(lines):
    rows = []
    for line in lines:
        if line.strip().startswith('|') and not re.match(r'\|\s*:?-+:?\s*\|', line):
            cols = [c.strip() for c in line.strip().split('|')[1:-1]]
            rows.append(cols)
    return rows

def add_table_from_rows(rows, header_bg="EAE3D9"):
    if not rows or len(rows) == 0:
        return
    num_cols = len(rows[0])
    tbl = doc.add_table(rows=len(rows), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Border xml
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="D8CEBF"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="D8CEBF"/><w:left w:val="single" w:sz="6" w:space="0" w:color="D8CEBF"/><w:right w:val="single" w:sz="6" w:space="0" w:color="D8CEBF"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="EAE3D9"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="EAE3D9"/></w:tblBorders>')
    tblPr.append(borders)
    
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            if c_idx >= num_cols:
                continue
            cell = tbl.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, bottom=80, left=110, right=110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            
            clean_val = val.replace('<br/>', '\n').replace('<br>', '\n').replace('&rarr;', '→').replace('&le;', '≤').replace('&ge;', '≥').replace('&times;', '×')
            
            if r_idx == 0:
                set_cell_background(cell, header_bg)
                run = p.add_run(clean_val)
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x2E, 0x1F, 0x18)
            else:
                if r_idx % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, "FAF7F2")
                tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', clean_val)
                for token in tokens:
                    if token.startswith('**') and token.endswith('**'):
                        run = p.add_run(token[2:-2])
                        run.font.bold = True
                    elif token.startswith('*') and token.endswith('*'):
                        run = p.add_run(token[1:-1])
                        run.font.italic = True
                    elif token.startswith('`') and token.endswith('`'):
                        run = p.add_run(token[1:-1])
                        run.font.name = 'Consolas'
                        run.font.size = Pt(8.5)
                    else:
                        run = p.add_run(token)
                    run.font.size = Pt(9.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def render_markdown_file(filepath):
    print(f"Rendering: {os.path.basename(filepath)}")
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_code = False
    code_lines = []
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle page breaks
        if '<div style="page-break-after: always;"></div>' in line or '<div class="page-break"></div>' in line:
            doc.add_page_break()
            i += 1
            continue

        # Handle Code Blocks
        if stripped.startswith('```'):
            if in_code:
                in_code = False
                caption = None
                if i + 1 < len(lines) and lines[i+1].strip().startswith('*Listing'):
                    caption = lines[i+1].strip().replace('*', '')
                    i += 1
                add_code_block(code_lines, caption)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
        
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if stripped.startswith('|') and '|' in stripped[1:]:
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                rows = parse_markdown_table(table_lines)
                add_table_from_rows(rows)
                table_lines = []

        # Handle Figure Images
        img_match = re.search(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            img_cap = img_match.group(1)
            img_rel = img_match.group(2)
            img_file = os.path.basename(img_rel)
            
            # Check for caption on next lines
            caption_line = None
            if i + 1 < len(lines) and lines[i+1].strip().startswith('**Figure'):
                caption_line = lines[i+1].strip().replace('**', '')
                i += 1
            elif i + 2 < len(lines) and lines[i+2].strip().startswith('**Figure'):
                caption_line = lines[i+2].strip().replace('**', '')
                i += 2
            add_figure_image(img_file, caption_line or img_cap)
            i += 1
            continue

        # Handle Headings
        if stripped.startswith('# '):
            add_chapter_title(stripped[2:])
            i += 1
            continue
        elif stripped.startswith('## '):
            add_heading_1(stripped[3:])
            i += 1
            continue
        elif stripped.startswith('### '):
            add_heading_2(stripped[4:])
            i += 1
            continue
        elif stripped.startswith('#### '):
            add_heading_3(stripped[5:])
            i += 1
            continue

        # Handle Bullets
        if stripped.startswith('- ') or stripped.startswith('* ') or re.match(r'^\d+\.\s+', stripped):
            bullet_text = re.sub(r'^(-|\*|\d+\.)\s+', '', stripped)
            add_bullet_item(bullet_text)
            i += 1
            continue

        # Handle horizontal divider
        if stripped in ['---', '***']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("_________________________________________________________________")
            run.font.color.rgb = RGBColor(0xD8, 0xCE, 0xBF)
            i += 1
            continue

        # Skip empty or html container lines
        if not stripped or stripped.startswith('<div') or stripped.startswith('</div>') or stripped.startswith('<br'):
            i += 1
            continue

        # Regular paragraph
        align = WD_ALIGN_PARAGRAPH.CENTER if ('align="center"' in line or 'text-align: center' in line) else WD_ALIGN_PARAGRAPH.JUSTIFY
        clean_text = re.sub(r'<.*?>', '', stripped)
        if clean_text:
            add_body_paragraph(clean_text, align)
        i += 1

    # End table if file ended inside table
    if in_table:
        rows = parse_markdown_table(table_lines)
        add_table_from_rows(rows)

# Compile all files sequentially
doc_order = [
    "front_matter.md",
    "ch1.md",
    "ch2.md",
    "ch3.md",
    "ch4.md",
    "ch5.md",
    "references.md",
    "glossary.md"
]

for idx, fname in enumerate(doc_order):
    fpath = os.path.join(base_dir, fname)
    if os.path.exists(fpath):
        render_markdown_file(fpath)
        if idx < len(doc_order) - 1:
            doc.add_page_break()
    else:
        print(f"Warning: {fpath} does not exist.")

doc.save(docx_path)
print(f"\n=======================================================")
print(f"SUCCESS: Final Blackbook document compiled successfully!")
print(f"Artifact Location: {docx_path}")
print(f"=======================================================")
